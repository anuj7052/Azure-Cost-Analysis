"""Invoice suggestions and reviewed monetary comparisons.

Extraction never posts a ledger entry: an admin reviews net amount, currency
and period first. Monetary arithmetic uses Decimal and never adds currencies.
"""
import asyncio
import io
import re
from decimal import Decimal
from urllib.parse import urlparse

import httpx
import pdfplumber
from openpyxl import Workbook, load_workbook
from core.config import settings

KINDS = ("customer", "distributor", "azure_usage")


def suggestions(text):
    out = {}
    patterns = {
        "invoice_number": r"(?:invoice\s*(?:number|no\.?|#))\s*[:\-]?\s*([A-Za-z0-9/\-]+)",
        "net_amount": r"(?:subtotal|sub.total|taxable\s*(?:value|amount)|total\s*(?:excluding|excl\.?)\s*tax)\s*[:\-]?\s*(?:[A-Z]{3}|[₹$€£])?\s*([0-9,]+\.\d{2})",
        "tax_amount": r"(?:total\s*tax|tax\s*amount)\s*[:\-]?\s*(?:[A-Z]{3}|[₹$€£])?\s*([0-9,]+\.\d{2})",
    }
    for field, pattern in patterns.items():
        match = re.search(pattern, text, re.I)
        if match:
            out[field] = match[1].replace(",", "") if field.endswith("amount") else match[1]
    currency = re.search(r"\b(INR|USD|EUR|GBP|AUD|CAD|AED|SGD)\b", text)
    if currency:
        out["currency"] = currency[1]
    elif "₹" in text:
        out["currency"] = "INR"
    # '$' alone is ambiguous. Do not guess a currency or interpret a gross total
    # as revenue excluding tax.
    return out


def read_text(content, extension):
    if extension == ".pdf":
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            return "\n".join((page.extract_text() or "")[:30000] for page in pdf.pages[:10])[:100000]
    if extension == ".csv":
        return content.decode("utf-8-sig", errors="replace")[:100000]
    if extension == ".xlsx":
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            lines = []
            for sheet in book.worksheets[:3]:
                for row in sheet.iter_rows(max_row=500, max_col=20, values_only=True):
                    lines.append(" ".join(str(v)[:500] for v in row if v is not None))
            return "\n".join(lines)[:100000]
        finally:
            book.close()
    return ""


async def extract_invoice(content, extension):
    try:
        text = await asyncio.to_thread(read_text, content, extension)
    except Exception:
        text = ""
    if text.strip():
        return {"method": "embedded_text", "suggestions": suggestions(text)}, "Review extracted suggestions against the original document. Only the first 10 PDF pages / 500 spreadsheet rows were inspected."
    if extension in (".csv", ".xlsx"):
        return {"method": "manual", "suggestions": {}}, "No readable invoice totals were extracted. Enter and verify the amounts manually."
    endpoint = settings.BILLING_OCR_ENDPOINT.rstrip("/")
    if not endpoint or not settings.BILLING_OCR_KEY:
        return {"method": "manual", "suggestions": {}}, "Document saved. Scanned-image OCR requires Azure Document Intelligence configuration; enter totals manually for now."
    parsed = urlparse(endpoint)
    if parsed.scheme != "https" or parsed.username or parsed.password or parsed.query or parsed.fragment or parsed.path not in {"", "/"} or not (parsed.hostname or "").endswith((".cognitiveservices.azure.com", ".api.cognitive.microsoft.com")):
        return {"method": "manual", "suggestions": {}}, "OCR endpoint configuration is invalid. Document saved; manual review is available."
    try:
        async with httpx.AsyncClient(timeout=25) as client:
            headers = {"Ocp-Apim-Subscription-Key": settings.BILLING_OCR_KEY, "Content-Type": "application/octet-stream"}
            response = await client.post(endpoint + "/documentintelligence/documentModels/prebuilt-invoice:analyze", params={"api-version": "2024-11-30", "pages": "1-10"}, headers=headers, content=content)
            response.raise_for_status()
            operation = response.headers.get("operation-location", "")
            op = urlparse(operation)
            if op.scheme != "https" or op.netloc != parsed.netloc:
                raise ValueError("Invalid operation location")
            for _ in range(15):
                response = await client.get(operation, headers={"Ocp-Apim-Subscription-Key": settings.BILLING_OCR_KEY})
                response.raise_for_status()
                result = response.json()
                if result.get("status") == "failed":
                    break
                if result.get("status") == "succeeded":
                    fields = ((result.get("analyzeResult", {}).get("documents") or [{}])[0]).get("fields", {})
                    values = {}
                    confidence = {}
                    for field, key in (("SubTotal", "net_amount"), ("TotalTax", "tax_amount"), ("InvoiceId", "invoice_number")):
                        entry = fields.get(field, {})
                        if field == "InvoiceId":
                            value = entry.get("valueString")
                        else:
                            money = entry.get("valueCurrency") or {}
                            value = money.get("amount")
                            if money.get("currencyCode"):
                                values["currency"] = money["currencyCode"]
                        if value is not None:
                            values[key] = str(value)
                            confidence[key] = entry.get("confidence")
                    return {"method": "azure_invoice_ocr", "suggestions": values, "confidence": confidence}, "OCR suggestions require review. Multiple invoices in one file must be uploaded separately."
                await asyncio.sleep(1)
    except (httpx.HTTPError, ValueError, KeyError, AttributeError):
        pass
    return {"method": "manual", "suggestions": {}}, "Document saved, but OCR did not complete. Download the original and enter the amounts manually."


def comparison(documents):
    groups = {}
    for row in documents:
        if row["status"] == "void":
            continue
        key = (row["period"], row["currency"] or "Unspecified")
        entry = groups.setdefault(key, {"period": key[0], "currency": key[1], "pending": 0, **{kind: None for kind in KINDS}})
        if row["status"] != "confirmed":
            entry["pending"] += 1
            continue
        kind = row["kind"]
        entry[kind] = (entry[kind] or Decimal(0)) + Decimal(row["net_amount"])
    result = []
    for entry in groups.values():
        revenue, cost, usage = (entry[k] for k in KINDS)
        profit = revenue - cost if revenue is not None and cost is not None else None
        variance = cost - usage if cost is not None and usage is not None else None
        pending_period = sum(row["status"] == "needs_review" and row["period"] == entry["period"] for row in documents)
        result.append({**{key: str(value) if isinstance(value, Decimal) else value for key, value in entry.items()},
                       "pending": pending_period,
                       "profit": str(profit) if profit is not None else None,
                       "distributor_vs_usage": str(variance) if variance is not None else None,
                       "margin_pct": str((profit / revenue * 100).quantize(Decimal("0.01"))) if profit is not None and revenue > 0 else None})
    return sorted(result, key=lambda item: (item["period"], item["currency"]), reverse=True)


def workbook(customer, documents, ledger_data=None):
    book = Workbook()
    sheet = book.active
    sheet.title = "Comparison"
    sheet.append(["Customer", "Period", "Currency", "Revenue net", "Distributor net", "Azure usage net", "Gross profit/loss", "Distributor less usage", "Margin %", "Pending review"])
    for row in comparison(documents):
        sheet.append([customer["name"], row["period"], row["currency"], *[float(row[k]) if row[k] is not None else None for k in ("customer", "distributor", "azure_usage", "profit", "distributor_vs_usage", "margin_pct")], row["pending"]])
    sheet = book.create_sheet("Documents")
    sheet.append(["Type", "Period", "Invoice number", "Currency", "Net", "Tax", "Status", "Original filename", "SharePoint URL"])
    for row in documents:
        sheet.append([row[k] for k in ("kind", "period", "invoice_number", "currency", "net_amount", "tax_amount", "status", "filename", "web_url")])
    if ledger_data:
        sheet = book.create_sheet('Monthly ledger')
        sheet.append(['Month', 'Currency', 'Customer recorded', 'Distributor recorded', 'Azure usage', 'Net profit/loss', 'Needs review'])
        for row in ledger_data['months']:
            sheet.append([row[k] for k in ('period', 'currency', 'customer_amount', 'distributor_amount', 'usage_amount', 'profit', 'review_count')])
        sheet = book.create_sheet('Workload entries')
        keys = ('workload', 'period', 'currency', 'invoice_number', 'customer_amount', 'distributor_amount', 'usage_amount', 'payment_status', 'amount_basis', 'status', 'notes', 'source_row', 'archived', 'workload_archived')
        sheet.append(list(keys))
        for row in ledger_data['entries']:
            sheet.append([row.get(k) for k in keys])
    for sheet in book:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for row in sheet:
            for cell in row:
                # Customer-controlled strings must never become Excel formulas.
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 24
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()
